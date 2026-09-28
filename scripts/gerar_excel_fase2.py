"""
gerar_excel_fase2.py - Gera analise-tecnica-<slug>.xlsx com 7 abas (Fase 1 + 2)

Uso: python scripts/gerar_excel_fase2.py --slug <slug> --dados <json>

O JSON de --dados tem as linhas REAIS de cada aba (nunca dado de exemplo
embutido aqui — erro 29 do Relatorio-Testes-4: um escritorio de advocacia
de Campinas ficou hardcoded neste arquivo e vazou pra planilha de qualquer
cliente novo). O agente monta esse JSON a partir de:
  - clusters.json (saida de skills/arquiteto-seo/scripts/clusterizar.py) —
    Inventario_Concorrente, Paginas_Concorrentes
  - projeto.md do cliente + pesquisa de concorrentes — Arquitetura_Sugerida,
    Plano_Construcao, Handoff_Fase3
  - raio-X tecnico da Fase 2 (o lider real, verificado) — RaioX_Tecnico,
    Schema_Concorrentes

Formato esperado (todas as 7 chaves obrigatorias, cada uma uma lista de
listas com os valores na ordem das colunas do cabecalho):

{
  "inventario_concorrente": [[dominio, url, trafego_mes, backlinks, ref_dominios, tipo, origem], ...],
  "paginas_concorrentes":   [[concorrente, url, kw_provavel, trafego_mes, backlinks, tipo, money_page_candidata], ...],
  "arquitetura_sugerida":   [[nivel, slug_sugerido, tipo, kw_principal, referencia_concorrente, observacao], ...],
  "plano_construcao":       [[prioridade, slug, kw_principal, volume_mes, sd, status, schema], ...],
  "raiox_tecnico":          [[sinal, status, detalhe, recomendacao], ...],
  "schema_concorrentes":    [[concorrente, pagina, schemas_detectados, json_ld, acao_recomendada], ...],
  "handoff_fase3":          [[pagina, kw_principal, volume_mes, estrutura_conteudo, schema, tom, status], ...]
}

Campo obrigatorio ausente ou aba vazia = falha com erro claro. Nunca cai em
valor padrao nem gera arquivo parcial.
"""
import argparse
import json
import os
import sys
from pathlib import Path

_PROJECT_DIR = Path(os.environ.get('CLAUDE_PROJECT_DIR', '.'))

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("FAIL: openpyxl nao instalado. Rode: pip install openpyxl")
    sys.exit(1)


ABAS = [
    ("inventario_concorrente", "Inventario_Concorrente",
     ["Dominio", "URL", "Trafego/mes", "Backlinks", "Ref. Dominios", "Tipo", "Origem"]),
    ("paginas_concorrentes", "Paginas_Concorrentes",
     ["Concorrente", "URL", "KW Provavel", "Trafego/mes", "Backlinks", "Tipo", "Money Page candidata?"]),
    ("arquitetura_sugerida", "Arquitetura_Sugerida",
     ["Nivel", "Slug sugerido", "Tipo", "KW principal", "Referencia concorrente", "Observacao"]),
    ("plano_construcao", "Plano_Construcao",
     ["Prioridade", "Slug", "KW principal", "Volume/mes", "SD", "Status", "Schema"]),
    ("raiox_tecnico", "RaioX_Tecnico",
     ["Sinal", "Status", "Detalhe", "Recomendacao para o cliente"]),
    ("schema_concorrentes", "Schema_Concorrentes",
     ["Concorrente", "Pagina", "Schemas detectados", "JSON-LD?", "Acao recomendada"]),
    ("handoff_fase3", "Handoff_Fase3",
     ["Pagina", "KW principal", "Volume/mes", "Estrutura de conteudo", "Schema", "Tom (OAB)", "Status"]),
]


def header_font():
    return Font(bold=True, color="FFFFFF", size=10)


def header_fill():
    return PatternFill("solid", fgColor="1F3864")


def alt_fill():
    return PatternFill("solid", fgColor="DCE6F1")


def thin_border():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)


def center_align():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def wrap_align():
    return Alignment(wrap_text=True, vertical="top")


def write_header(ws, cols):
    ws.append(cols)
    for i in range(1, len(cols) + 1):
        c = ws.cell(ws.max_row, i)
        c.font = header_font()
        c.fill = header_fill()
        c.alignment = center_align()
        c.border = thin_border()
    ws.row_dimensions[ws.max_row].height = 28


def write_data(ws, vals, alt=False):
    ws.append(vals)
    for i in range(1, len(vals) + 1):
        c = ws.cell(ws.max_row, i)
        if alt:
            c.fill = alt_fill()
        c.alignment = wrap_align()
        c.border = thin_border()


def auto_width(ws, min_w=12, max_w=55):
    for col in ws.columns:
        w = min_w
        for cell in col:
            try:
                if cell.value:
                    w = max(w, min(max_w, len(str(cell.value)) + 2))
            except Exception:
                pass
        ws.column_dimensions[get_column_letter(col[0].column)].width = w


def carregar_dados(caminho):
    p = Path(caminho)
    if not p.exists():
        print(f"FAIL: arquivo de dados nao encontrado: {caminho}")
        print("Monte o JSON com as linhas reais de cada aba antes de rodar este script "
              "(ver docstring deste arquivo) — nunca rode sem --dados.")
        sys.exit(1)
    dados = json.loads(p.read_text(encoding="utf-8"))
    faltando = []
    vazias = []
    for chave, _, _ in ABAS:
        if chave not in dados:
            faltando.append(chave)
        elif not dados[chave]:
            vazias.append(chave)
    if faltando:
        print(f"FAIL: chaves ausentes no JSON de dados: {', '.join(faltando)}")
        sys.exit(1)
    if vazias:
        print(f"FAIL: abas sem nenhuma linha (vazias nao sao permitidas): {', '.join(vazias)}")
        sys.exit(1)
    return dados


def gerar(slug, caminho_dados):
    dados = carregar_dados(caminho_dados)

    pasta = _PROJECT_DIR / f"projetos/{slug}"
    saida = pasta / f"analise-tecnica-{slug}.xlsx"

    wb = openpyxl.Workbook()
    primeira = True
    for chave, titulo, cols in ABAS:
        ws = wb.active if primeira else wb.create_sheet(titulo)
        if primeira:
            ws.title = titulo
            primeira = False
        ws.freeze_panes = "A2"
        write_header(ws, cols)
        for i, row in enumerate(dados[chave]):
            if len(row) != len(cols):
                print(f"FAIL: aba {titulo}, linha {i + 1} tem {len(row)} valores, "
                      f"esperado {len(cols)} ({', '.join(cols)})")
                sys.exit(1)
            write_data(ws, list(row), alt=(i % 2 == 1))
        auto_width(ws)

    pasta.mkdir(parents=True, exist_ok=True)
    wb.save(str(saida))
    print(f"PASS - Excel gerado: {saida} ({wb.sheetnames})")
    return 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--slug", required=True)
    parser.add_argument("--dados", required=True,
                         help="JSON com as linhas reais das 7 abas (ver docstring deste arquivo)")
    args = parser.parse_args()
    sys.exit(gerar(args.slug, args.dados))
