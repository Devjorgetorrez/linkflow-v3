"""
criar_excel_fase1.py - Gera analise-concorrentes-<slug>.xlsx a partir do projeto.md
Uso: python scripts/criar_excel_fase1.py --slug <slug>
Requer: pip install openpyxl
"""

import argparse
import os
import re
from pathlib import Path

_PROJECT_DIR = Path(os.environ.get('CLAUDE_PROJECT_DIR', '.'))

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERRO: openpyxl nao instalado. Rode: pip install openpyxl")
    raise SystemExit(1)


def parse_table(conteudo, secao):
    match = re.search(rf"## {secao}.*?\n(.*?)(?=\n##|\Z)", conteudo, re.DOTALL | re.IGNORECASE)
    if not match:
        return [], []
    linhas = match.group(1).split("\n")
    headers, rows = [], []
    for linha in linhas:
        if not linha.strip().startswith("|"):
            continue
        cols = [c.strip() for c in linha.split("|") if c.strip()]
        if not cols:
            continue
        if re.match(r"^[-:]+$", cols[0]):
            continue
        if not headers:
            headers = cols
        else:
            rows.append(cols)
    return headers, rows


def criar_excel(slug):
    caminho_md = _PROJECT_DIR / f"projetos/{slug}/projeto.md"
    if not caminho_md.exists():
        print(f"ERRO: {caminho_md} nao encontrado")
        raise SystemExit(1)

    conteudo = caminho_md.read_text(encoding="utf-8")

    wb = openpyxl.Workbook()

    # ── ABA 1: Páginas dos Concorrentes ──────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Paginas Concorrentes"

    h_pag, rows_pag = parse_table(conteudo, "Paginas dos Concorrentes")

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    trans_fill  = PatternFill("solid", fgColor="E2EFDA")
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, h in enumerate(h_pag, 1):
        cell = ws1.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for row_idx, row in enumerate(rows_pag, 2):
        transacional = row[2].lower().startswith("sim") if len(row) > 2 else False
        for col_idx, val in enumerate(row, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True)
            if transacional:
                cell.fill = trans_fill

    for col_idx in range(1, len(h_pag) + 1):
        ws1.column_dimensions[get_column_letter(col_idx)].width = 30

    ws1.freeze_panes = "A2"

    # ── ABA 2: Money Pages ───────────────────────────────────────────────────
    ws2 = wb.create_sheet("Money Pages")

    h_mp, rows_mp = parse_table(conteudo, "Money Pages")

    for col_idx, h in enumerate(h_mp, 1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    alta_fill = PatternFill("solid", fgColor="FFF2CC")
    for row_idx, row in enumerate(rows_mp, 2):
        alta = any("ALTA" in v.upper() for v in row)
        for col_idx, val in enumerate(row, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border
            if alta:
                cell.fill = alta_fill

    for col_idx in range(1, len(h_mp) + 1):
        ws2.column_dimensions[get_column_letter(col_idx)].width = 25

    ws2.freeze_panes = "A2"

    # ── Salvar ───────────────────────────────────────────────────────────────
    saida = _PROJECT_DIR / f"projetos/{slug}/analise-concorrentes-{slug}.xlsx"
    wb.save(saida)
    print(f"OK: {saida} criado com {len(rows_pag)} paginas e {len(rows_mp)} Money Pages")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--slug", required=True)
    args = parser.parse_args()
    criar_excel(args.slug)
